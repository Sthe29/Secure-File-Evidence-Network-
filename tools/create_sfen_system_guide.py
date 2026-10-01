from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(r"C:\Users\sithe\Downloads\SFEN system\SFEN System Guide.docx")

BLUE = "155EEF"
NAVY = "14213D"
PALE_BLUE = "EDF4FF"
LIGHT_GRAY = "D9D9D9"
WHITE = "FFFFFF"
BLACK = "000000"

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:fill'), fill)
    tc_pr.append(shd)

def border(cell, color=LIGHT_GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tc_pr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '6')
        el.set(qn('w:color'), color)
        borders.append(el)

def set_cell_margins(cell, top=110, start=110, bottom=110, end=110):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{m}'))
        if node is None:
            node = OxmlElement(f'w:{m}')
            tcMar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def add_page_field(paragraph):
    run = paragraph.add_run('Page ')
    fldChar1 = OxmlElement('w:fldChar'); fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText'); instrText.set(qn('xml:space'), 'preserve'); instrText.text = 'PAGE'
    fldChar2 = OxmlElement('w:fldChar'); fldChar2.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1); run._r.append(instrText); run._r.append(fldChar2)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_body(doc, text):
    p = doc.add_paragraph(style='Body Text')
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.12
    p.add_run(text)
    return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p

def add_number(doc, text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)
    return p

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.7)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.78)
section.right_margin = Inches(0.78)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos')
styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
styles['Normal'].font.size = Pt(10.5)
for name, size in [('Title', 24), ('Heading 1', 16), ('Heading 2', 12.5)]:
    style = styles[name]
    style.font.name = 'Aptos Display' if name != 'Heading 2' else 'Aptos'
    style._element.rPr.rFonts.set(qn('w:ascii'), style.font.name)
    style._element.rPr.rFonts.set(qn('w:hAnsi'), style.font.name)
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)

if 'Body Text' not in styles:
    styles.add_style('Body Text', WD_STYLE_TYPE.PARAGRAPH)
styles['Body Text'].font.name = 'Aptos'
styles['Body Text'].font.size = Pt(10.5)

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
run = header.add_run('SFEN  |  Secure File and Evidence Network')
run.font.name = 'Aptos'; run.font.size = Pt(8.5); run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run('SFEN System Guide  |  ')
footer_run.font.size = Pt(8.5); footer_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
add_page_field(footer)

title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.add_run('SFEN System Guide')
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = subtitle.add_run('Secure File and Evidence Network')
r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = RGBColor(0, 0, 0)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('Role responsibilities, case workflow, evidence access, audit logging, and database records')
r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
doc.add_paragraph()

add_heading(doc, 'Purpose and Scope', 1)
add_body(doc, 'SFEN is a role based police case management prototype. It connects public incident reporting with station intake, official CAS registration, detective investigation, station commander oversight, administrative control, and a PostgreSQL database record. The system is designed to keep the original online report, the official police case, the evidence, the movement of the docket, and the activity history connected.')
add_body(doc, 'The system uses one registered detective in the current prototype. Cases are therefore automatically allocated to Detective Inspector David Khumalo after the officer completes official case registration. This removes unnecessary manual assignment controls while keeping the flow clear for testing.')

add_heading(doc, 'Main Roles in SFEN', 1)
table = doc.add_table(rows=1, cols=3)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Table Grid'
headers = ['Role', 'Main responsibility', 'What the role can do']
for i, text in enumerate(headers):
    cell = table.rows[0].cells[i]
    cell.text = text
    shade(cell, NAVY); border(cell); set_cell_margins(cell)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for run in cell.paragraphs[0].runs:
        run.font.bold = True; run.font.color.rgb = RGBColor(255,255,255); run.font.size = Pt(9.5)
set_repeat_table_header(table.rows[0])
rows = [
    ('Citizen', 'Report an incident and follow its progress.', 'Creates an account using a 13 digit ID number and 10 digit mobile number, submits an online report, uploads supporting evidence, views the evidence later, tracks the online reference and CAS case updates, and receives notifications.'),
    ('CSC Police Officer', 'Verify the report at the police station and open the official case.', 'Uses one Review and Register Case flow, checks the citizen identity, records the formal statement, records station evidence, confirms the declaration, creates the official CAS number, and moves the docket into the investigation workflow.'),
    ('Detective', 'Investigate the official CAS docket.', 'Views the case file and citizen evidence in Documents, acknowledges docket custody, adds investigation diary entries, adds detective documents, changes status with a reason, records exhibits, and follows commander directives.'),
    ('Station Commander', 'Provide supervision and accountability.', 'Views transferred or assisted cases, reads citizen evidence and authorised documents in Documents, reviews progress, issues instructions, records supervisory reviews, monitors docket movement, and sees the audit trail.'),
    ('System Administrator', 'Manage controlled access and monitor system activity.', 'Creates personnel accounts, provides temporary passwords, manages users and roles, views activity logs, and verifies system records. New personnel must set their own permanent password at first login.')
]
for row_i, values in enumerate(rows):
    cells = table.add_row().cells
    for i, text in enumerate(values):
        cells[i].text = text
        shade(cells[i], PALE_BLUE if row_i % 2 == 0 else WHITE)
        border(cells[i]); set_cell_margins(cells[i])
        cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cells[i].paragraphs[0].runs:
            run.font.size = Pt(8.7)

add_heading(doc, 'System Flow From Start to Finish', 1)
add_body(doc, 'The case journey is designed as one connected record. The online reference does not replace a CAS case number. The online reference begins the report, while the official CAS number is created by a police officer after in person verification and registration.')
flow = [
    'Citizen account creation. The citizen enters a full name, email address, 13 digit South African ID number, 10 digit mobile number, and password. The system validates these details before the account is created.',
    'Online incident report. The citizen selects the incident type, enters the date and time, provides a location, writes a description, adds people or property details, and may attach photos, documents, audio, or video evidence.',
    'Online reference number. After submission, SFEN creates an online report reference such as SFEN RPT number. The citizen can open My Reports to view what was submitted and verify uploaded evidence.',
    'Officer review. The CSC officer sees the report in the station report queue. The officer selects Review and Register Case, which first opens the complete report details, location information, narrative, and uploaded evidence.',
    'Station verification and statement. When the citizen attends the station, the officer verifies the ID, captures the formal statement in the system, records any evidence received at the station, and confirms the sworn declaration.',
    'CAS registration. Only after the verification conditions are complete can the officer register the official case. SFEN creates the CAS number and preserves a link to the original online report.',
    'Automatic detective allocation. In the current system one detective is registered, so the newly registered docket is automatically routed to Detective Inspector David Khumalo. The detective acknowledges docket custody before actively progressing the investigation.',
    'Investigation. The detective works from the docket workspace, reviews the citizen uploaded evidence under Documents, records investigation diary entries, attaches further documents, logs exhibits, and changes the case status with a recorded reason.',
    'Commander oversight. The commander can review the case, evidence, documents, movements, and diary progress. The commander may issue instructions and record supervisory review outcomes. If the case is transferred to the commander, that handover remains visible in the docket history.',
    'Citizen updates and close out. The citizen sees the online reference, CAS number, docket timeline, and notifications. When the case progresses, the record remains available rather than being deleted, which protects traceability and audit accountability.'
]
for item in flow:
    add_number(doc, item)

add_heading(doc, 'Citizen Portal', 1)
add_body(doc, 'The citizen portal is the public entry point. A citizen must sign in or create an account before submitting a report. The dashboard gives direct access to reporting, records, profile information, notifications, and the case journey.')
add_bullet(doc, 'Report Incident uses a staged form for incident classification, location, narrative, supporting information, attachments, and final review.')
add_bullet(doc, 'Evidence uploaded by the citizen is stored with the report. New uploads can be opened later by the citizen, authorised officer, detective, and station commander.')
add_bullet(doc, 'Immediate attention reports can be marked as urgent. Officers receive the alert with the reporter details, description, and mapped location where supplied.')
add_bullet(doc, 'My Reports lets the citizen verify submitted information and retain the online reference number. My Cases shows the official CAS case after station registration.')

add_heading(doc, 'CSC Officer Workflow', 1)
add_body(doc, 'The CSC officer is responsible for converting a preliminary report into an official police case. The officer does not create a CAS number merely because an online form exists. The officer must perform the in station verification process first.')
add_bullet(doc, 'Open the report through Review and Register Case, not through separate review and registration buttons.')
add_bullet(doc, 'Review the original incident information, citizen contact details, location, description, and uploaded evidence.')
add_bullet(doc, 'Verify the citizen ID, capture the formal statement, record optional evidence identifying details such as firearm serial numbers, and confirm the consent or declaration.')
add_bullet(doc, 'Register the CAS docket only when the required identity, statement, and declaration checks are complete.')
add_bullet(doc, 'The officer audit trail combines officer actions with docket movements so the officer can show what they did and when.')

add_heading(doc, 'Detective Investigation Workflow', 1)
add_body(doc, 'The detective works on the official CAS docket after it has been registered and allocated. The detective does not need an Assign Case control because the current station setup has one detective. The docket workspace provides the operational case record.')
add_bullet(doc, 'Acknowledge receipt of the docket to record the chain of custody.')
add_bullet(doc, 'Open Documents to see citizen uploaded photos, documents, video, or audio, as well as documents attached during investigation.')
add_bullet(doc, 'Use the investigation diary for chronological notes, witness interviews, evidence analysis, court updates, and actions taken.')
add_bullet(doc, 'When changing the case status, provide a reason. The reason is retained in the diary and audit history.')
add_bullet(doc, 'Record exhibits and evidence references so physical and digital items can be traced during the investigation.')

add_heading(doc, 'Station Commander Workflow', 1)
add_body(doc, 'The station commander provides oversight rather than performing the frontline intake. The commander can inspect a docket that is under supervision, transferred to the commander, or requires intervention.')
add_bullet(doc, 'Open Documents to view citizen uploaded evidence and authorised case documents when assisting with a case or after a transfer.')
add_bullet(doc, 'Review the detective progress, diary entries, audit trail, and docket movement history.')
add_bullet(doc, 'Record supervisory reviews and issue specific directions to the detective.')
add_bullet(doc, 'Monitor the current custodian of the docket and acknowledge or return a transferred docket where appropriate.')

add_heading(doc, 'System Administrator Workflow', 1)
add_body(doc, 'The system administrator manages account access and system accountability. Police personnel do not create their own accounts. The administrator creates the account and receives a temporary password from the system.')
add_bullet(doc, 'Create a personnel account with the officer name, personnel number, departmental email, contact number, rank, role, and assigned station.')
add_bullet(doc, 'Give the temporary password to the new personnel member through an approved channel.')
add_bullet(doc, 'At first login the personnel member must choose their own permanent password before continuing to the system.')
add_bullet(doc, 'Use System Activity to inspect recorded actions such as reports submitted, cases registered, password events, transfers, and status changes.')

add_heading(doc, 'Evidence and Documents', 1)
add_body(doc, 'Evidence is treated as part of the case record. A citizen can upload supporting media with an online report. Newly uploaded files retain a viewable copy in the local prototype so authorised users can confirm the correct material was supplied.')
add_bullet(doc, 'Citizen: views the original submitted file in My Reports.')
add_bullet(doc, 'Officer: views the file during report review and station registration.')
add_bullet(doc, 'Detective: views the file under the official case Documents tab and can use it during investigation.')
add_bullet(doc, 'Commander: views the file under Documents during oversight, assistance, or transfer review.')
add_bullet(doc, 'Older prototype uploads that saved only file names and metadata cannot be previewed. New uploads contain the viewable copy.')

add_heading(doc, 'Audit Logging and Docket Movement', 1)
add_body(doc, 'SFEN records important system actions in activity logs. This supports accountability because a police case should not rely only on a current screen state. The history records the actor, action, reference number, date and time, and linked record where applicable.')
add_bullet(doc, 'Examples include citizen account creation, online report submission, urgent report submission, case registration, officer review, detective status updates, diary entries, personnel account creation, and password setup.')
add_bullet(doc, 'Docket movement records show who sent the docket, who should receive it, the reason for the transfer, the date and time, and whether the recipient acknowledged receipt.')
add_bullet(doc, 'Records are retained for traceability. The preferred case outcome is an appropriate status or closure record, not deletion of police history.')

add_heading(doc, 'Database and Data Integrity', 1)
add_body(doc, 'SFEN uses PostgreSQL through Prisma. The database provides the persistent data layer for accounts, reports, official cases, audits, transfers, investigation diary entries, and supervisory reviews.')
db_table = doc.add_table(rows=1, cols=2)
db_table.alignment = WD_TABLE_ALIGNMENT.CENTER
db_table.style = 'Table Grid'
for i, text in enumerate(['Table', 'Purpose']):
    cell = db_table.rows[0].cells[i]; cell.text = text; shade(cell, NAVY); border(cell); set_cell_margins(cell)
    for run in cell.paragraphs[0].runs:
        run.font.bold = True; run.font.color.rgb = RGBColor(255,255,255); run.font.size = Pt(9.5)
set_repeat_table_header(db_table.rows[0])
db_rows = [
    ('User', 'Citizen and authorised personnel accounts, roles, identity details, and authentication fields.'),
    ('IncidentReport', 'Online and walk in intake reports, location, narrative, attachments, and report status.'),
    ('Case', 'Official CAS dockets linked to a single source report.'),
    ('AuditEntry', 'Recorded system activity associated with a user, case, report, or other entity.'),
    ('DocketTransfer', 'Docket handover, current custody, recipient, acknowledgement, and reason.'),
    ('InvestigationDiaryEntry', 'Detective investigation actions, outcomes, status related reasons, and next actions.'),
    ('SupervisoryReview', 'Commander review outcomes, notes, directions, and review dates.'),
    ('_prisma_migrations', 'Technical migration history used by the database framework.')
]
for row_i, values in enumerate(db_rows):
    cells = db_table.add_row().cells
    for i, text in enumerate(values):
        cells[i].text = text; shade(cells[i], PALE_BLUE if row_i % 2 == 0 else WHITE); border(cells[i]); set_cell_margins(cells[i])
        for run in cells[i].paragraphs[0].runs: run.font.size = Pt(9)

add_heading(doc, 'Validation and Access Controls', 1)
add_bullet(doc, 'Citizen ID numbers are limited to 13 digits and citizen mobile numbers to 10 digits.')
add_bullet(doc, 'A case cannot be registered until the officer has verified identity, captured the statement, and certified the declaration.')
add_bullet(doc, 'Personnel accounts are created by an administrator and use a temporary password before a permanent first login password is set.')
add_bullet(doc, 'API requests require a valid authenticated session. Expired sessions return the user to sign in before a sensitive action such as case registration can continue.')
add_bullet(doc, 'Access is role based: citizens see their own reports, while operational users see records needed for their assigned responsibilities.')

add_heading(doc, 'Recommended Demonstration Order', 1)
add_body(doc, 'For a presentation, demonstrate one complete journey rather than opening unrelated pages. This order makes the connection between the user interface, the roles, activity log, and database easy to understand.')
for item in [
    'Create or sign in as a citizen and submit a short report with a photo or PDF.',
    'Open My Reports and preview the uploaded evidence and online reference number.',
    'Sign in as the CSC officer, select Review and Register Case, review the report, and complete the station verification conditions.',
    'Register the CAS docket and show that it is linked to the original online report.',
    'Sign in as the detective, open Documents, preview the citizen evidence, acknowledge custody, and add an investigation update with a reason.',
    'Sign in as the station commander, open the same case, review Documents and the audit trail, and record or explain a supervisory instruction.',
    'Sign in as administrator and show System Activity and the PostgreSQL tables in pgAdmin.'
]:
    add_number(doc, item)

add_heading(doc, 'Conclusion', 1)
add_body(doc, 'SFEN gives one connected path from a citizen report to an official CAS docket and investigation record. Each role has a focused responsibility, while evidence, movements, updates, and activity history remain connected to the case. This supports clearer communication with citizens, better operational visibility for police personnel, and a database record that can be demonstrated in pgAdmin.')

doc.save(OUT)
print(OUT)
